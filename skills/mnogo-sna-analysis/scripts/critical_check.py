"""Read-only independent audit of the Mnogo Sna workbook/HTML workflow.

No network requests or source writes. --node executes only an explicitly supplied,
task-authored dashboard script with a test DOM; it is not a browser or a security sandbox.
"""
import argparse
import datetime as dt
import hashlib
import itertools
import json
import math
import posixpath
import re
import subprocess
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal as D
from pathlib import Path

import openpyxl


class CheckError(Exception):
    pass


def require(condition, message):
    if not condition:
        raise CheckError(message)


def numeric(value):
    require(isinstance(value, (int, float, D)) and not isinstance(value, bool), 'Пустое или нечисловое значение')
    result = D(str(value))
    require(result.is_finite(), 'Нечисловой бесконечный показатель')
    return result


def near(a, b, tolerance='0.0000001'):
    return abs(numeric(a) - numeric(b)) <= D(tolerance)


def normalize(value):
    text = unicodedata.normalize('NFKC', str(value or '')).strip().casefold()
    text = re.sub(r'^заказ\s+(?:клиента|покупателя)\s+', '', text)
    return re.sub(r'\s+', '', text).replace('–', '-').replace('—', '-')


def address_key(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def visible_rows(path, sheet, data_only=True):
    book = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    try:
        require(sheet in book.sheetnames, f'Нет листа {sheet}: {Path(path).name}')
        tab = book[sheet]
        require(tab.sheet_state == 'visible', f'Запрещено читать скрытый лист {sheet}')
        return list(tab.values)
    finally:
        book.close()


def cell(rows, row, col):
    return rows[row-1][col-1] if row <= len(rows) and col <= len(rows[row-1]) else None


def sums(rows):
    return {k: float(sum((numeric(x[k]) for x in rows), D('0'))) for k in ('qty', 'amount')}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sheet_structure(path, sheet):
    """Read only the selected worksheet's XML; never load hidden worksheet parts."""
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    relns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(path) as z:
        workbook=ET.fromstring(z.read('xl/workbook.xml'))
        tab=next(x for x in workbook.find('s:sheets',ns) if x.get('name')==sheet)
        require(tab.get('state','visible')=='visible','Запрещено читать структуру скрытого листа')
        rid=tab.get('{'+relns+'}id')
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target=next(x.get('Target') for x in rels if x.get('Id')==rid)
        part=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        xml=ET.fromstring(z.read(part))
        rows={int(x.get('r')):{'hidden':x.get('hidden','0').lower() in ('1','true'),'outline':int(x.get('outlineLevel','0')),'collapsed':x.get('collapsed','0').lower() in ('1','true'),'height':float(x.get('ht','0'))} for x in xml.findall('s:sheetData/s:row',ns)}
        merges={x.get('ref') for x in xml.findall('s:mergeCells/s:mergeCell',ns)}
        return rows,merges


def check_sources(m):
    require(m.get('source_hashes'), 'Не заданы исходные хеши: версию источников проверить нельзя')
    for path, expected in m['source_hashes'].items():
        require(Path(path).is_file(), 'Нет исходного файла: '+Path(path).name)
        require(sha(path) == expected, 'Изменилась версия источника: '+Path(path).name)
    return {'sources': len(m['source_hashes']), 'unchanged': True}


def check_price(m):
    c = m['price']
    original = visible_rows(c['source'], c['sheet'])
    result = visible_rows(c['output'], c['sheet'])
    formulas = visible_rows(c['output'], c['sheet'], False)
    model_row, header_row = c['model_row'], c['output_header_row']
    models = [(col, name) for col, name in enumerate(original[model_row-1], 1) if col > c['size_col'] and name]
    size_rows = [(r, str(cell(original, r, c['size_col'])).strip()) for r in range(1, len(original)+1)
                 if re.fullmatch(r'\d+(?:[.,]\d+)?[xх*×]\d+(?:[.,]\d+)?', str(cell(original, r, c['size_col'])).strip(), re.I)]
    require(models and size_rows, 'Не распознаны модели или размеры прайса')
    costbook = openpyxl.load_workbook(c['cost_source'], read_only=True, data_only=True)
    checked = 0
    try:
        for i, (col, model) in enumerate(models):
            cost_col = c['output_first_cost_col'] + i*c['output_block_cols']
            require(cell(result, model_row, cost_col) == model, f'Сдвинуто название модели {model}')
            require([cell(result, header_row, cost_col+j) for j in range(3)] == ['Себестоимость','Доля себестоимости','Цена'], f'Неверные заголовки {model}')
            require(model in costbook.sheetnames, f'Нет точного совпадения модели {model}')
            tab = costbook[model]
            require(tab.sheet_state == 'visible', f'Совпадение находится на скрытом листе: {model}')
            costs = list(tab.values)
            headers = [(r, k) for r, row in enumerate(costs[:3], 1) for k, v in enumerate(row, 1) if v == c['cost_header']]
            require(len(headers) == 1, f'Заголовок себестоимости не уникален: {model}')
            _, sc = headers[0]
            for row_no, size in size_rows:
                width, length = [D(x.replace(',','.')) for x in re.split('[xх*×]',size,flags=re.I)]
                matches = []
                for rr in range(1, len(costs)+1):
                    if cell(costs, rr, sc-4) != model:
                        continue
                    w, l, value = cell(costs, rr, sc-3), cell(costs, rr, sc-2), cell(costs, rr, sc)
                    if isinstance(w,(int,float)) and isinstance(l,(int,float)) and near(numeric(w)*100,width) and near(numeric(l)*100,length):
                        matches.append(value)
                require(len(matches) == 1, f'Неединственное совпадение {model}, {size}')
                expected_cost = numeric(matches[0])
                price = numeric(cell(original, row_no, col))
                require(price > 0, f'Неположительная цена {model}, {size}')
                require(near(cell(result,row_no,cost_col),expected_cost),f'Ошибка себестоимости {model}, {size}')
                require(near(cell(result,row_no,cost_col+2),price),f'Изменена исходная цена {model}, {size}')
                require(near(cell(result,row_no,cost_col+1),expected_cost/price),f'Ошибка кешированной доли {model}, {size}')
                a = openpyxl.utils.get_column_letter(cost_col)+str(row_no)
                b = openpyxl.utils.get_column_letter(cost_col+2)+str(row_no)
                formula = str(cell(formulas,row_no,cost_col+1) or '').replace('$','').replace(' ','')
                require(formula == f'={a}/{b}', f'Формула доли ссылается не на свою цену: {model}, {size}')
                checked += 1
    finally:
        costbook.close()
    # Unchanged auxiliary visible sheets are audited without opening hidden content.
    orig = openpyxl.load_workbook(c['source'], read_only=True, data_only=True)
    out = openpyxl.load_workbook(c['output'], read_only=True, data_only=True)
    try:
        require(orig.sheetnames == out.sheetnames, 'Изменён состав листов прайса')
        for name in orig.sheetnames:
            require(orig[name].sheet_state == out[name].sheet_state, 'Изменена видимость листа прайса')
            if name != c['sheet'] and orig[name].sheet_state == 'visible':
                require(list(orig[name].values) == list(out[name].values), 'Изменён дополнительный лист: '+name)
    finally:
        orig.close(); out.close()
    return {'models':len(models),'sizes':len(size_rows),'costs_prices_and_formulas':checked}


def parse_shipments(m):
    c=m['shipments']
    rows=visible_rows(c['source'],c['sheet'])
    header=rows[c['header_row']-1]
    qi,ai=header.index(c['qty_header']),header.index(c['amount_header'])
    date=None;order=None;current_day=None;current_order=None
    days=[];orders=[];items=[];missing=[];others=[];grand=None
    for r,row in enumerate(rows,1):
        text=str(row[0] or '').strip()
        if re.fullmatch(r'\d{2}\.\d{2}\.\d{4}',text):
            date=dt.datetime.strptime(text,'%d.%m.%Y').date().isoformat();order=None
            current_day={'row':r,'qty':row[qi],'amount':row[ai],'items':[]};days.append(current_day)
        elif (match:=re.match(r'^Заказ\s+(?:клиента|покупателя)\s+(.+?)\s+от\s+',text,re.I)):
            require(date is not None,'Заказ без даты отгрузки')
            order=match[1].strip();current_order={'row':r,'qty':row[qi],'amount':row[ai],'items':[]};orders.append(current_order)
        elif text.casefold()=='итого':
            grand={'qty':row[qi],'amount':row[ai]}
        elif r>c['header_row']+2 and text:
            require(date and order,f'Товар вне иерархии, строка {r}')
            if not re.match(r'^Матрас\b',text,re.I):
                others.append(r);continue
            base={'row':r,'order':order,'date':date,'month':date[:7],'product':text}
            if row[qi] is None or row[ai] is None:
                missing.append({**base,'qty':row[qi],'amount':row[ai]});continue
            size=re.findall(r'(\d+(?:[.,]\d+)?)\s*[*xх×]\s*(\d+(?:[.,]\d+)?)',text,re.I)
            require(len(size)==1,f'Размер требует уточнения, строка {r}')
            width,length=[D(x.replace(',','.')) for x in size[0]]
            qty,amount=numeric(row[qi]),numeric(row[ai])
            require(qty>=0 and amount>=0,f'Возврат/отрицательная строка требует правила, строка {r}')
            record={**base,'qty':float(qty),'amount':float(amount),'width':float(width),'length':float(length)}
            items.append(record);current_day['items'].append(record);current_order['items'].append(record)
    require(grand is not None,'Нет общего итога отгрузок')
    require(sums(items)==grand,'Товарные строки не равны общему итогу отгрузок')
    for block in days+orders:
        if block['qty'] is not None and block['amount'] is not None:
            require(sums(block['items'])=={k:block[k] for k in ('qty','amount')},f'Не сходится промежуточный итог, строка {block["row"]}')
    return items,missing,others,rows


def check_dashboard(m):
    items,missing,others,source_rows=parse_shipments(m)
    text=Path(m['dashboard']).read_text(encoding='utf-8')
    match=re.search(r'<script\s+id="payload"\s+type="application/json">(.*?)</script>',text,re.S)
    require(match is not None,'Нет доступных данных payload в конечном HTML')
    data=json.loads(match[1]);actual=data['items']
    require(len({x['row'] for x in actual})==len(actual),'Задвоены товарные строки HTML')
    by_row={x['row']:x for x in actual}
    require(set(by_row)=={x['row'] for x in items},'Потеряны или добавлены товарные строки HTML')
    index={}
    for path in m['orders']:
        rows=visible_rows(path,m['shipments']['sheet'])
        header=rows[0];ni,ai=header.index('Номер'),header.index('Адрес доставки')
        for rn,row in enumerate(rows[1:],2):
            if row[ni]:index.setdefault(normalize(row[ni]),[]).append((address_key(row[ai]),Path(path).name,rn,openpyxl.utils.get_column_letter(ai+1)))
    known_map=m['city_regions']
    def infer_region(address):
        s=address.casefold().replace('ё','е')
        if re.search(r'\bмосква\b|\bмосковская\s+обл',s):return 'Москва и Московская область'
        if 'санкт-петербург' in s or re.search(r'\bленинградская\s+обл',s):return 'Санкт-Петербург и Ленинградская область'
        explicit=[region for region in set(known_map.values()) if region.casefold() in s]
        city=[region for name,region in known_map.items() if re.search(r'\b'+re.escape(name)+r'\b',address,re.I)]
        candidates=set(explicit+city)
        require(len(candidates)==1,'География не подтверждена словарём/адресом; нужна ручная проверка')
        return candidates.pop()
    for expected in items:
        row=by_row[expected['row']]
        for k in ('order','date','month','product','qty','amount','width','length'):
            require(row[k]==expected[k],f'HTML не соответствует отгрузкам: строка {expected["row"]}, поле {k}')
        group='До 120 см включительно' if expected['width']<=120 else 'Свыше 120 см'
        require(row['width_group']==group,'Неверная группа ширины, строка '+str(expected['row']))
        matches=index.get(normalize(expected['order']),[])
        addresses={x[0] for x in matches if x[0]}
        if len(addresses)==1:
            address=addresses.pop()
            require(row['address']==address,'Подтянут не тот адрес, строка '+str(expected['row']))
            require(row['region']==infer_region(address),'Неверный регион, строка '+str(expected['row']))
            refs=' / '.join(f'{file}: {m["shipments"]["sheet"]}!{col}{rn}' for _,file,rn,col in matches)
            require(row['order_source_ref']==refs,'Неверная ссылка на источник адреса')
        else:
            require(row['region']=='Требует уточнения','Неясный адрес ошибочно распределён')
    require({x['row'] for x in data['anomalies']}=={x['row'] for x in missing},'Потеряны строки без показателей')
    for expected in missing:
        row=next(x for x in data['anomalies'] if x['row']==expected['row'])
        for k,v in expected.items():
            if k!='month':require(row[k]==v,'Неверно описана строка без показателей')
    require(len(data['other_products'])==len(others),'Не показаны посторонние товары')
    require(sums(actual)==sums(items),'Итог HTML не равен источнику')
    require(all(data['meta'][k]==v for k,v in sums(items).items()),'Итоги meta не равны исходным товарным строкам')
    require(len(data['monthly'])==len({x['month'] for x in actual}) and {x['month'] for x in data['monthly']}=={x['month'] for x in actual},'Потерян или задвоен месяц сводки')
    require(len(data['summary'])==len({x['region'] for x in actual}) and {x['region'] for x in data['summary']}=={x['region'] for x in actual},'Потерян или задвоен регион сводки')
    for month in {x['month'] for x in items}:
        expected=sums([x for x in items if x['month']==month]);stored=next(x for x in data['monthly'] if x['month']==month)
        require(all(stored[k]==v for k,v in expected.items()),'Неверный месяц '+month)
    for stored in data['summary']:
        region_rows=[x for x in actual if x['region']==stored['region']];expected=sums(region_rows);grand=sums(actual)
        require(all(stored[k]==v for k,v in expected.items()),'Не сходится регион '+stored['region'])
        for k in ('qty','amount'):require(near(stored[k+'_share'],numeric(expected[k])/numeric(grand[k])),'Неверный знаменатель доли')
    # Source report and grouped rows in the enriched workbook remain intact.
    s=visible_rows(m['shipments']['source'],m['shipments']['sheet'],False)
    o=visible_rows(m['shipments']['output'],m['shipments']['sheet'],False)
    for rn,row in enumerate(s,1):
        for col,value in enumerate(row,1):require(value==cell(o,rn,col),'Изменена исходная ячейка '+openpyxl.utils.get_column_letter(col)+str(rn))
    source_dims,source_merges=sheet_structure(m['shipments']['source'],m['shipments']['sheet'])
    output_dims,output_merges=sheet_structure(m['shipments']['output'],m['shipments']['sheet'])
    for rn,dim in source_dims.items():
        other=output_dims.get(rn,{})
        require(all(dim[k]==other.get(k) for k in ('hidden','outline','collapsed')),'Изменена группировка строки '+str(rn))
        require(abs(dim['height']-other.get('height',0))<0.2,'Изменена высота строки '+str(rn))
    require(source_merges<=output_merges,'Утеряны объединения отчёта')
    for row in data['enriched']:
        rn=row['row']
        for col,key in [(8,'order'),(9,'address'),(10,'region'),(11,'width'),(12,'width_group'),(13,'status'),(14,'order_source_ref')]:
            require((cell(o,rn,col) or '')==(row[key] or ''),'Ошибка добавленной колонки '+str(rn))
    return data,text,{'item_rows':len(items),'missing_metric_rows':len(missing),'other_product_rows':len(others),**sums(items),'addresses_matched':sum(x['region']!='Требует уточнения' for x in actual)}


def check_tariffs(m,data):
    p=data['shipping'];approved=m['shipping_parameters']
    for key,value in approved.items():require(p[key]==value,'Неверный параметр ПЭК: '+key)
    dims=[D(str(x)) for x in p['dimensions_m']]
    require(math.prod(dims)==D(str(p['volume_m3'])),'Объём не равен произведению точных габаритов')
    require(p['destination_mode']=='Доставить по адресу','Получение на терминале вместо адреса')
    quotes={q['city']:q for q in p['quotes']}
    require(len(quotes)==len(p['quotes']),'Повторы тарифов города')
    for q in quotes.values():
        require(near(numeric(q['auto'])+numeric(q['delivery']),q['total'],'0.005'),'Не сходятся компоненты тарифа '+q['city'])
        require(near(numeric(q['total'])+numeric(q['insurance'])+numeric(q['insuranceOrganization']),q['rawTotal'],'0.005'),'Полный тариф включает неучтённые услуги '+q['city'])
    costs=D('0');qty=D('0');cities=set()
    for x in data['items']:
        if x['region'] in p['excluded_regions']:
            require(x['shipping_unit'] is None and x['shipping_total'] is None,'Исключённый регион включён в доставку')
            continue
        candidates=[city for city in quotes if re.search(r'\b'+re.escape(city)+r'\b',x['address'],re.I)]
        require(len(candidates)==1,'Реальный населённый пункт не определён: строка '+str(x['row']))
        city=candidates[0];require(x['shipping_city']==city,'Населённый пункт подменён центром региона')
        require(near(x['shipping_unit'],quotes[city]['total'],'0.005'),'Подставлен не тот тариф города')
        expected=numeric(x['qty'])*numeric(quotes[city]['total'])
        require(near(x['shipping_total'],expected,'0.005'),'Ошибка количества × тарифа')
        costs+=expected;qty+=numeric(x['qty']);cities.add(city)
    require(near(costs,p['total_estimate'],'0.005') and qty==numeric(p['eligible_qty']),'Не сходится общий расход ПЭК')
    require(cities==set(quotes),'Лишние или пропущенные города в тарифах')
    for c in p['cities']:
        rows=[x for x in data['items'] if x['region']==c['region'] and x['shipping_city']==c['city']]
        require(numeric(c['qty'])==sum((numeric(x['qty']) for x in rows),D('0')),'Неверное количество города')
        require(near(c['shipping_total'],sum((numeric(x['shipping_total']) for x in rows),D('0')),'0.005'),'Неверный расход города')
    return {'cities':len(cities),'eligible_qty':float(qty),'estimate':float(costs),'provider_live_rechecked':False}


NODE_TEST = r'''
const vm=require('node:vm'),assert=require('node:assert/strict');let input='';
process.stdin.setEncoding('utf8');process.stdin.on('data',s=>input+=s);process.stdin.on('end',()=>{
const {html,cases}=JSON.parse(input),payload=html.match(/<script id="payload" type="application\/json">([\s\S]*?)<\/script>/)[1],code=html.match(/<script>\s*([\s\S]*?)<\/script>/)[1];
const ids=new Set([...html.matchAll(/\bid="([^"]+)"/g)].map(x=>x[1])),nodes=new Map;
class Element{constructor(id){this.value=['month','region','width'].includes(id)?'all':'';this.options=[];this.textContent=id==='payload'?payload:'';this.innerHTML='';this.classList={toggle(){}}}add(o){this.options.push(o)}addEventListener(){}setAttribute(){}get selectedOptions(){return [{textContent:this.options.find(o=>o.value===this.value)?.textContent||this.value}]}}
const document={getElementById(id){assert(ids.has(id),'HTML lacks '+id);if(!nodes.has(id))nodes.set(id,new Element(id));return nodes.get(id)}};
const context={document,window:{},Intl,URL,Blob,setTimeout,Option:function(t,v){this.textContent=t;this.value=v}};vm.createContext(context);vm.runInContext(code,context,{timeout:10000});
for(const c of cases){for(const [id,v]of [['month',c.month],['region',c.region],['width',c.width]])document.getElementById(id).value=v;context.window.dashboard.render();const s=context.window.dashboard.getState();assert.deepEqual(JSON.parse(JSON.stringify(s.total)),c.total);assert.deepEqual(JSON.parse(JSON.stringify(s.denom)),c.denom);assert.equal(s.shippingQty,c.shippingQty);assert(Math.abs(s.shippingTotal-c.shippingTotal)<0.001);assert.equal(s.rows.length,c.rows)}
console.log(JSON.stringify({cases:cases.length,engine:'Node test DOM',visual_browser_check:false}));});
'''


def check_js(node,html,data):
    cases=[]
    for month,width,region in itertools.product(['all']+sorted({x['month'] for x in data['items']}),['all','narrow','wide'],['all']+sorted({x['region'] for x in data['items']})):
        base=[x for x in data['items'] if (month=='all' or x['month']==month) and (width=='all' or (x['width']<=120 if width=='narrow' else x['width']>120))]
        rows=[x for x in base if region=='all' or x['region']==region]
        shipping=[x for x in rows if x['shipping_unit'] is not None]
        cases.append({'month':month,'width':width,'region':region,'total':sums(rows),'denom':sums(base),'rows':len(rows),'shippingQty':sum(x['qty'] for x in shipping),'shippingTotal':float(sum((numeric(x['shipping_total']) for x in shipping),D('0')))})
    result=subprocess.run([node,'-e',NODE_TEST],input=json.dumps({'html':html,'cases':cases}),text=True,encoding='utf-8',capture_output=True,timeout=60)
    require(result.returncode==0,'JavaScript пересчёт не совпал с Python: '+result.stderr[-1200:])
    return json.loads(result.stdout)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest',type=Path,required=True)
    parser.add_argument('--node',help='Явно разрешённый runtime для теста собственного HTML; не применять к чужому непроверенному коду')
    parser.add_argument('--expect-snapshot',action='store_true')
    parser.add_argument('--report',type=Path,help='Создать новый отчёт JSON; существующий файл не перезаписывается')
    args=parser.parse_args();m=json.loads(args.manifest.read_text(encoding='utf-8'))
    report={'status':'NUMERIC_PASS','checks':[],'warnings':[],'manual_checks':m.get('manual_checks',{})}
    def run(name,fn):
        try:
            value=fn();report['checks'].append({'name':name,'status':'PASS','result':value});return value
        except Exception as exc:
            report['checks'].append({'name':name,'status':'FAIL','error':str(exc)});report['status']='FAIL';return None
    run('source_versions',lambda:check_sources(m))
    if m.get('price'):run('price_cost_share',lambda:check_price(m))
    parsed=None
    if m.get('dashboard'):
        holder={}
        def parse():
            data,html,result=check_dashboard(m);holder.update(data=data,html=html);return result
        run('shipments_orders_dashboard',parse)
        if holder:
            data,html=holder['data'],holder['html']
            if data.get('shipping'):run('pec_recorded_tariffs_and_totals',lambda:check_tariffs(m,data))
            if args.node and data.get('shipping'):run('javascript_filter_recalculation',lambda:check_js(args.node,html,data))
            elif data.get('shipping'):report['warnings'].append('JavaScript-пересчёт не исполнен: --node не указан.')
            if args.expect_snapshot:
                def snapshot():
                    for key,value in m['snapshot_totals'].items():
                        actual={'qty':data['meta']['qty'],'amount':data['meta']['amount'],'item_rows':len(data['items']),'narrow_qty':sum(x['qty'] for x in data['items'] if x['width']<=120),'wide_qty':sum(x['qty'] for x in data['items'] if x['width']>120),'shipping_qty':data['shipping']['eligible_qty'],'shipping_estimate':data['shipping']['total_estimate']}[key]
                        require(near(actual,value,'0.005'),'Изменён исторический контроль '+key)
                    return {'snapshot':m['snapshot_date']}
                run('historical_snapshot_only',snapshot)
    if not m.get('currency_confirmed'):
        report['warnings'].append('Валюта выручки не подтверждена источником/Ириной; значок ₽ в HTML — неподтверждённое допущение.')
    report['warnings'].extend([
        'Записанные тарифы проверены математически, но скрипт не запрашивает свежие цены ПЭК. Дата и параметры должны быть проверены отдельно.',
        'Скрипт не доказывает открытие Excel без восстановления или визуальную проверку браузером. Статусы ручных проверок — записи манифеста, не независимое подтверждение.',
        'NUMERIC_PASS не означает полную управленческую приёмку или подтверждение фактических расходов/оплаты доставки.'
    ])
    output=json.dumps(report,ensure_ascii=False,indent=2)
    print(output)
    if args.report:
        require(not args.report.exists(),'Отчёт уже существует; перезапись запрещена')
        args.report.write_text(output,encoding='utf-8')
    return 2 if report['status']=='FAIL' else 0


if __name__=='__main__':
    sys.exit(main())
